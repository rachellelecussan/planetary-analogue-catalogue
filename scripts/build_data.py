"""Convert the data-entry spreadsheet into the files the website reads.

Usage:
    python scripts/build_data.py [path/to/catalogue.xlsx]

Reads the Sites, References and Vocabularies tabs and writes:
    data/sites.geojson   one feature per site, with its references attached
    data/vocab.json      the controlled vocabularies (used by filters and the form)

Coordinates are rounded according to coordinate_precision (and sensitive sites
are never published more precisely than about 10 km), so exact locations of
fragile sites never reach the public website.
"""
import json
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
SRC = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "data-source" / "catalogue.xlsx"

# Columns that are internal to the workbook and never published
PRIVATE = {"notes_for_reviewer", "reference_count", "required_missing", "entry_check"}
MULTI = {"alt_names", "analogue_for_other", "analogue_type_other", "environment_type_other",
         "rock_types", "key_minerals", "hazards", "contributors"}
DECIMALS = {"Exact": 4, "Approximately 1 km": 2, "Approximately 10 km": 1, "Region only": 0}


def clean(v):
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def split(v):
    return [p.strip() for p in str(v).split(";") if p.strip()] if v else []


def main():
    wb = load_workbook(SRC, data_only=True)
    ws = wb["Sites"]
    headers = [str(ws.cell(row=2, column=c).value or "").replace(" *", "").strip()
               for c in range(1, ws.max_column + 1)]

    refs = {}
    rf = wb["References"]
    for r in range(4, rf.max_row + 1):
        sid = clean(rf.cell(row=r, column=1).value)
        cite = clean(rf.cell(row=r, column=2).value)
        if not sid or not cite or cite.startswith("[AUTHOR"):
            continue
        link = clean(rf.cell(row=r, column=3).value)
        notes = clean(rf.cell(row=r, column=6).value) or ""
        refs.setdefault(sid, []).append({
            "citation": cite,
            "doi": link if link and not str(link).startswith("http") else None,
            "url": link if link and str(link).startswith("http") else (f"https://doi.org/{link}" if link else None),
            "supports": split(clean(rf.cell(row=r, column=4).value)),
            "peer_reviewed": notes.startswith("Peer-reviewed"),
            "unverified": "NOT VERIFIED" in notes,
        })

    def table(name):
        """Rows of an optional tab as dicts (header on row 2, data from row 3)."""
        if name not in wb.sheetnames:
            return []
        t = wb[name]
        heads = [t.cell(row=2, column=c).value for c in range(1, t.max_column + 1)]
        out = []
        for r in range(3, t.max_row + 1):
            row = {h: clean(t.cell(row=r, column=c + 1).value) for c, h in enumerate(heads) if h}
            if row.get("site_id"):
                out.append(row)
        return out

    comparisons = {}
    for c in table("Comparisons"):
        if c.get("latitude") is None or c.get("longitude") is None:
            continue
        sources = []
        for kind in ("pairing", "coordinate"):
            if c.get(f"{kind}_source_url"):
                sources.append({"title": c.get(f"{kind}_source_title") or c[f"{kind}_source_url"], "url": c[f"{kind}_source_url"]})
        comparisons[c["site_id"]] = {
            "body": c.get("body"), "area": c.get("area_name"), "lat": float(c["latitude"]), "lon": float(c["longitude"]),
            "view_km": float(c.get("view_km") or 200), "why": c.get("why_compared"),
            "confidence": c.get("confidence"), "sources": sources,
        }

    photos = {}
    photo_dir = ROOT / "assets" / "photos"
    for p in table("Photos"):
        if not p.get("file") or not (photo_dir / p["file"]).exists():
            continue
        photos.setdefault(p["site_id"], []).append({
            "file": p["file"], "caption": p.get("caption"), "author": p.get("author"),
            "licence": p.get("licence"), "licence_url": p.get("licence_url"), "source": p.get("source_page"),
        })

    features, problems = [], []
    for r in range(4, ws.max_row + 1):
        row = {h: clean(ws.cell(row=r, column=c + 1).value) for c, h in enumerate(headers) if h}
        if not row.get("id") or not row.get("name"):
            continue
        lat, lon = row.pop("latitude", None), row.pop("longitude", None)
        if lat is None or lon is None:
            problems.append(f"{row['id']}: missing coordinates, skipped")
            continue
        prec = row.get("coordinate_precision") or "Approximately 10 km"
        dp = DECIMALS.get(prec, 1)
        if row.get("sensitive_site") == "Yes":
            dp = min(dp, 1)
        props = {}
        for k, v in row.items():
            if k in PRIVATE or v is None:
                continue
            props[k] = split(v) if k in MULTI else v
        for k in ("access_level", "hazard_rating"):
            if k in props:
                props[k] = int(props[k])
        props["references"] = refs.get(row["id"], [])
        if row["id"] in comparisons:
            props["compare"] = comparisons[row["id"]]
        props["photos"] = photos.get(row["id"], [])
        if not props["references"]:
            problems.append(f"{row['id']}: no references")
        features.append({
            "type": "Feature",
            "geometry": {"type": "Point", "coordinates": [round(float(lon), dp), round(float(lat), dp)]},
            "properties": props,
        })

    out = {"type": "FeatureCollection",
           "generated": datetime.now().strftime("%Y-%m-%d"),
           "features": features}
    (ROOT / "data" / "sites.geojson").write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")

    vocab = {}
    vs = wb["Vocabularies"]
    for c in range(1, vs.max_column + 1):
        key = vs.cell(row=1, column=c).value
        if not key:
            continue
        vals = []
        for r in range(2, vs.max_row + 1):
            v = vs.cell(row=r, column=c).value
            if v is None:
                break
            vals.append(v)
        vocab[key] = vals
    (ROOT / "data" / "vocab.json").write_text(json.dumps(vocab, ensure_ascii=False, indent=1), encoding="utf-8")

    print(f"Wrote {len(features)} sites, {sum(len(f['properties']['references']) for f in features)} references")
    for p in problems:
        print("WARNING:", p)


if __name__ == "__main__":
    main()
